"""Text extraction from uploaded documents."""

from __future__ import annotations

import io
from dataclasses import dataclass, field

SUPPORTED_MIME_TYPES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/plain": "txt",
    "text/markdown": "md",
    "text/x-markdown": "md",
}

SUPPORTED_EXTENSIONS = {".pdf": "pdf", ".docx": "docx", ".txt": "txt", ".md": "md"}


class UnsupportedDocument(Exception):
    pass


@dataclass(slots=True)
class ParsedPage:
    number: int
    text: str


@dataclass(slots=True)
class ParsedDocument:
    text: str
    pages: list[ParsedPage] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)


def detect_kind(filename: str, mime_type: str | None) -> str:
    if mime_type and mime_type.split(";")[0].strip() in SUPPORTED_MIME_TYPES:
        return SUPPORTED_MIME_TYPES[mime_type.split(";")[0].strip()]

    lowered = (filename or "").lower()
    for suffix, kind in SUPPORTED_EXTENSIONS.items():
        if lowered.endswith(suffix):
            return kind

    raise UnsupportedDocument(
        f"Unsupported document '{filename}' ({mime_type}). Supported: PDF, DOCX, TXT, MD."
    )


def parse_document(data: bytes, filename: str, mime_type: str | None = None) -> ParsedDocument:
    kind = detect_kind(filename, mime_type)
    if kind == "pdf":
        return _parse_pdf(data)
    if kind == "docx":
        return _parse_docx(data)
    return _parse_text(data)


def _parse_pdf(data: bytes) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(ParsedPage(number=number, text=text))

    return ParsedDocument(
        text="\n\n".join(page.text for page in pages),
        pages=pages,
        metadata={"page_count": len(reader.pages)},
    )


def _parse_docx(data: bytes) -> ParsedDocument:
    import docx

    document = docx.Document(io.BytesIO(data))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]

    # Tables often hold the facts people later ask about.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))

    text = "\n\n".join(paragraphs)
    return ParsedDocument(text=text, pages=[ParsedPage(number=1, text=text)] if text else [])


def _parse_text(data: bytes) -> ParsedDocument:
    text = data.decode("utf-8", errors="replace").strip()
    return ParsedDocument(text=text, pages=[ParsedPage(number=1, text=text)] if text else [])
